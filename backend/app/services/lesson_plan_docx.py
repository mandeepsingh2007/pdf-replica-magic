"""Fill the supplied lesson-plan template and export it as PDF."""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
import threading
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "lesson-plan-template.docx"
_WORD = threading.Lock()

LEVELS = (
    ("remembering", "1. REMEMBERING (Knowledge Recall)", "Identify", "Remembering"),
    ("understanding", "2. UNDERSTANDING (Comprehension)", "Explain", "Understanding"),
    ("applying", "3. APPLYING (Use in Real-Life Context)", "Use", "Applying"),
    ("analyzing", "4. ANALYZING (Thinking & Differentiating)", "Differentiate", "Analyzing"),
    ("evaluating", "5. EVALUATING (Judgement & Expression)", "Express", "Evaluating"),
    ("creating", "6. CREATING (Creativity & Expression)", "Create", "Creating"),
)


def _lines(values) -> list[str]:
    cleaned = []
    for value in values or []:
        text = " ".join(str(value).split())
        if text:
            cleaned.append(text)
    return cleaned


def _ordinal_suffix(level: int) -> str:
    return {1: "st", 2: "nd", 3: "rd"}.get(level, "th")


def _texts(element):
    return element.findall(".//" + qn("w:t"))


def _set_text(node, value: str) -> None:
    node.text = value
    if value[:1].isspace() or value[-1:].isspace():
        node.set(qn("xml:space"), "preserve")


def _replace_exact(element, old: str, new: str) -> None:
    for node in _texts(element):
        if node.text == old:
            _set_text(node, new)
            return


def _clone(element):
    return deepcopy(element)


def _append(cell, element) -> None:
    cell._tc.append(element)


def _with_text(proto, text: str):
    element = _clone(proto)
    nodes = _texts(element)
    if nodes:
        _set_text(nodes[0], text)
    return element


def _objective(proto, verb: str, rest: str, level_name: str):
    element = _clone(proto)
    nodes = _texts(element)
    cleaned = " ".join(rest.split())
    if cleaned.lower().startswith(verb.lower()):
        cleaned = cleaned[len(verb) :].strip()
    _set_text(nodes[0], verb)
    _set_text(nodes[1], f" {cleaned} " if cleaned else " ")
    _set_text(nodes[2], f"({level_name})")
    return element


def _bold_run(element):
    for run in element.findall(qn("w:r")):
        if run.find(qn("w:rPr") + "/" + qn("w:b")) is not None and run.find(qn("w:t")) is not None:
            return run
    return None


def _plain_run(element):
    for run in element.findall(qn("w:r")):
        if run.find(qn("w:t")) is not None and run.find(qn("w:rPr") + "/" + qn("w:b")) is None:
            return run
    runs = [run for run in element.findall(qn("w:r")) if run.find(qn("w:t")) is not None]
    return runs[0]


def _conclusion(proto, text: str, phrases: list[str]):
    element = _clone(proto)
    plain = _clone(_plain_run(element))
    bold_run = _bold_run(element)
    bold = _clone(bold_run if bold_run is not None else plain)
    for run in list(element.findall(qn("w:r"))):
        element.remove(run)
    remaining = " ".join(text.split())
    unused = [phrase.strip() for phrase in phrases if phrase and phrase.strip() in remaining]
    while remaining:
        found = None
        at = len(remaining)
        for phrase in unused:
            index = remaining.find(phrase)
            if index != -1 and index < at:
                found = phrase
                at = index
        if found is None:
            run = _clone(plain)
            _set_text(run.find(qn("w:t")), remaining)
            element.append(run)
            break
        if at:
            run = _clone(plain)
            _set_text(run.find(qn("w:t")), remaining[:at])
            element.append(run)
        run = _clone(bold)
        _set_text(run.find(qn("w:t")), found)
        element.append(run)
        remaining = remaining[at + len(found) :]
    return element


def _clear_paragraphs(cell) -> None:
    for child in list(cell._tc):
        if child.tag == qn("w:p"):
            cell._tc.remove(child)


def _fill(doc: Document, *, class_level: int, subject: str, chapter: str, plan: dict) -> None:
    cell = doc.tables[0].rows[1].cells[1]
    source = list(cell.paragraphs)
    title = source[0]._element
    subtitle = source[1]._element
    meta = source[2]._element
    chapter_line = source[3]._element
    spacer = source[4]._element
    objectives_heading = source[5]._element
    intro = source[6]._element
    objective = source[7]._element
    rule = source[13]._element
    section = source[14]._element
    bloom = source[16]._element
    label = source[17]._element
    bullet = source[18]._element
    closing = source[98]._element
    closing_body = source[99]._element

    meta_line = _clone(meta)
    _replace_exact(meta_line, "1", str(class_level))
    _replace_exact(meta_line, "st", _ordinal_suffix(class_level))
    _replace_exact(meta_line, "English", subject)

    chapter_el = _clone(chapter_line)
    _replace_exact(chapter_el, "A Special Gift", chapter)

    _clear_paragraphs(cell)
    _append(cell, _clone(title))
    _append(cell, _clone(subtitle))
    _append(cell, meta_line)
    _append(cell, chapter_el)
    _append(cell, _clone(spacer))
    _append(cell, _clone(objectives_heading))
    _append(cell, _clone(intro))
    for key, _heading, verb, level_name in LEVELS:
        _append(cell, _objective(objective, verb, plan[key].get("objective") or "", level_name))

    _append(cell, _clone(rule))
    _append(cell, _clone(section))
    _append(cell, _clone(rule))
    for key, heading, _verb, _level_name in LEVELS:
        body = plan[key]
        _append(cell, _with_text(bloom, heading))
        _append(cell, _with_text(label, "Activity:"))
        for line in _lines(body.get("activity")):
            _append(cell, _with_text(bullet, line))
        _append(cell, _with_text(label, "Student Task:"))
        for line in _lines(body.get("student_task")):
            _append(cell, _with_text(bullet, line))
        _append(cell, _with_text(label, "Assessment:"))
        for line in _lines(body.get("assessment")):
            _append(cell, _with_text(bullet, line))
        _append(cell, _clone(rule))

    _append(cell, _with_text(bloom, "ASSESSMENT & HOMEWORK"))
    _append(cell, _with_text(label, "Assessment:"))
    for line in _lines(plan.get("assessment")):
        _append(cell, _with_text(bullet, line))
    _append(cell, _with_text(label, "Homework:"))
    for line in _lines(plan.get("homework")):
        _append(cell, _with_text(bullet, line))
    _append(cell, _clone(spacer))
    _append(cell, _clone(closing))
    _append(cell, _conclusion(closing_body, plan.get("conclusion") or "", plan.get("conclusion_phrases") or []))


def _docx_to_pdf_word(docx_path: str, pdf_path: str) -> None:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(os.path.abspath(docx_path), False, False, False)
        document.ExportAsFixedFormat(os.path.abspath(pdf_path), 17)
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def _docx_to_pdf_libreoffice(docx_path: str, pdf_path: str) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice is required to create the lesson plan PDF on this server.")
    outdir = os.path.dirname(os.path.abspath(pdf_path))
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--nolockcheck",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf",
        "--outdir",
        outdir,
        os.path.abspath(docx_path),
    ]
    result = subprocess.run(
        cmd,
        check=False,
        capture_output=True,
        timeout=180,
        env={**os.environ, "SAL_USE_VCLPLUGIN": "svp", "HOME": "/tmp"},
    )
    produced = os.path.join(outdir, f"{Path(docx_path).stem}.pdf")
    if result.returncode != 0 or not os.path.isfile(produced):
        detail = (result.stderr or result.stdout or b"").decode("utf-8", "replace")[-400:]
        raise RuntimeError(f"Lesson plan PDF could not be created. {detail}".strip())
    if os.path.abspath(produced) != os.path.abspath(pdf_path):
        os.replace(produced, pdf_path)


def _docx_to_pdf(docx_path: str, pdf_path: str) -> None:
    if os.name == "nt":
        _docx_to_pdf_word(docx_path, pdf_path)
        return
    _docx_to_pdf_libreoffice(docx_path, pdf_path)


def build_lesson_plan_pdf(
    *,
    class_level: int,
    subject: str,
    chapter: str,
    plan: dict,
) -> bytes:
    if not TEMPLATE.is_file():
        raise RuntimeError("Lesson plan template is missing.")
    doc = Document(str(TEMPLATE))
    _fill(doc, class_level=class_level, subject=subject, chapter=chapter, plan=plan)
    header = doc.tables[0].rows[0]._tr.get_or_add_trPr()
    if header.find(qn("w:tblHeader")) is None:
        header.append(OxmlElement("w:tblHeader"))
    folder = tempfile.mkdtemp(prefix="lesson-plan-")
    docx_path = os.path.join(folder, "plan.docx")
    pdf_path = os.path.join(folder, "plan.pdf")
    doc.save(docx_path)
    with _WORD:
        _docx_to_pdf(docx_path, pdf_path)
    data = Path(pdf_path).read_bytes()
    if not data.startswith(b"%PDF"):
        raise RuntimeError("Lesson plan PDF could not be created.")
    return data
